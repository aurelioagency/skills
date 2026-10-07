#!/usr/bin/env python3
"""Scrapea soloduenos.com con Scrapling y vuelca los avisos a un Excel (fuente de verdad).

Uso:
  python scrape.py login                         # abre el navegador del perfil para iniciar sesion a mano (una vez)
  python scrape.py run --excel RUTA.xlsx [--max-pages 3] [--ops venta alquiler alquiler_temporario]
                       [--zona "capital federal"] [--sin-contacto] [--headless]

El Excel se actualiza por link: lo nuevo se agrega, lo existente no se pisa
(las columnas de gestion quedan intactas). Los campos de contacto se completan
solo si estan vacios.
"""
import argparse, re, sys, time, unicodedata, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from scrapling.fetchers import DynamicSession

BASE = "https://www.soloduenos.com"
FUENTE = "soloduenos.com"
PROFILE = Path.home() / ".soloduenos-profile"
STATE = PROFILE / "state.json"  # cookies + localStorage de la sesion iniciada a mano

COLUMNAS = [
    "ID", "Fuente", "Link publicación", "Fecha de captura",
    "Nombre del dueño", "Teléfono / WhatsApp", "Email", "Otro contacto",
    "Tipo", "Operación", "Zona", "Precio y moneda", "Características",
    "Detalle breve", "Días publicado",
    "Estado de contacto", "Fecha de contacto", "Notas",
]
GESTION = {"Estado de contacto", "Fecha de contacto", "Notas"}


def norm(s):
    s = unicodedata.normalize("NFD", s or "")
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


# ---------- Excel ----------
def abrir_excel(path):
    path = Path(path)
    if path.exists():
        wb = load_workbook(path)
        ws = wb["Leads"]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"
        ws.append(COLUMNAS)
        wb.create_sheet("Corridas").append(["Fecha", "Operaciones", "Zona", "Avisos vistos", "Nuevos", "Con teléfono"])
        ws.freeze_panes = "A2"
        for col, w in zip("ABCDEFGHIJKLMNOPQR", [10, 15, 50, 16, 22, 20, 24, 20, 14, 12, 28, 18, 36, 60, 10, 18, 16, 30]):
            ws.column_dimensions[col].width = w
    return wb, ws


def indice_por_link(ws):
    hdr = [c.value for c in ws[1]]
    ci = hdr.index("Link publicación") + 1
    return hdr, {ws.cell(r, ci).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, ci).value}


# ---------- Scrapling ----------
def listar_links(session, op, max_pages):
    links, vistos = [], set()
    for page in range(1, max_pages + 1):
        r = session.fetch(f"{BASE}/buscar?operation_type={op}&page={page}", wait_selector="a[href*='/propiedades/']")
        nuevos = [h for h in r.css("a[href*='/propiedades/']::attr(href)").getall()]
        nuevos = [h if h.startswith("http") else BASE + h for h in nuevos]
        nuevos = [h for h in dict.fromkeys(nuevos) if h not in vistos]
        if not nuevos:
            break
        vistos.update(nuevos)
        links += nuevos
        time.sleep(1)
    return links


def accion_detalle(con_contacto, out):
    """page_action: si se pide contacto, abre el boton de WhatsApp y junta lo que aparezca."""
    def _accion(page):
        if not con_contacto:
            return
        try:
            popups = []
            page.context.on("page", lambda p: popups.append(p))
            page.locator("button:has-text('Contactar por WhatsApp')").first.click(timeout=4000)
            page.wait_for_timeout(2500)
            out["login"] = bool(re.search(r"Ingres[aá].{0,40}cre[aá] tu cuenta", page.inner_text("body")))
            hrefs = page.eval_on_selector_all(
                "a[href*='wa.me'],a[href*='whatsapp'],a[href^='tel:'],a[href^='mailto:']", "e=>e.map(x=>x.href)")
            out["contacto"] = " | ".join(hrefs)
            if popups:
                out["contacto"] += " | " + popups[0].url
                popups[0].close()
        except Exception:
            pass
    return _accion


def parsear(url, body):
    lines = [l.strip() for l in body.splitlines() if l.strip()]
    d = {"url": url}
    m = re.search(r"/propiedades/(venta|alquiler-temporario|alquiler)-", url)
    d["operacion"] = {"venta": "Venta", "alquiler": "Alquiler", "alquiler-temporario": "Alquiler temporario"}.get(m.group(1), "") if m else ""
    m = re.search(r"C[oó]digo\s*[·•\-]?\s*#?(\d+)", body)
    d["id"] = m.group(1) if m else url.rsplit("-", 5)[-5:][0]
    d["titulo"], d["zona"], d["tipo"] = "", "", ""
    for i, l in enumerate(lines):
        if re.match(r"C[oó]digo", l):
            d["titulo"] = lines[i - 1] if i else ""
            d["zona"] = lines[i + 1] if i + 1 < len(lines) else ""
            break
    m = re.search(r"/propiedades/[a-z\-]+?-(departamento|casa|ph|oficina|local|terreno|campo|cochera|galpon|propiedad)-", url)
    d["tipo"] = ("PH" if m and m.group(1) == "ph" else m.group(1).capitalize()) if m else ""
    m = re.search(r"((?:US\$|\$|USD)\s*[\d\.\,]+)", body)
    d["precio"] = m.group(1).replace(" ", " ") if m else ""
    car = []
    for l in lines:
        if re.search(r"\b(m² (cubiertos|totales)|ambientes?|dormitorios?|baños?|cocheras?|a[nñ]os)\b", l) and len(l) < 40:
            car.append(l)
    d["caracteristicas"] = ", ".join(dict.fromkeys(car))
    desc = ""
    if "Descripción" in lines:
        i = lines.index("Descripción")
        desc = lines[i + 1] if i + 1 < len(lines) else ""
    d["descripcion"] = desc
    d["nombre"] = ""
    for i, l in enumerate(lines):
        if l.startswith("Dueño/a") and i + 1 < len(lines) and lines[i + 1] not in ("Dueño", "Dueño/a"):
            d["nombre"] = lines[i + 1]
    return d


def detalle_breve(d):
    base = " ".join(x for x in [d["tipo"], d["caracteristicas"].split(",")[0] if d["caracteristicas"] else "", "en", d["zona"].split(",")[0]] if x)
    extra = d["descripcion"].split(".")[0][:140]
    return f"{base}. {extra}".strip(". ") if extra else base


def parsear_contacto(txt, popup):
    blob = f"{txt} {popup}"
    tel = ""
    m = re.search(r"wa\.me/(\d+)|phone=(\d+)|tel:(\+?\d+)", blob)
    if m:
        tel = next(g for g in m.groups() if g)
    mail = re.search(r"mailto:([^\s|?]+)", blob)
    return tel, (mail.group(1) if mail else "")


# ---------- comandos ----------
def cmd_login(_):
    print(f"Se abre un navegador. Inicia sesion en soloduenos.com y despues cerra la ventana. Sesion en {STATE}")
    from playwright.sync_api import sync_playwright
    PROFILE.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(str(PROFILE / "chromium"), headless=False)
        (ctx.pages[0] if ctx.pages else ctx.new_page()).goto(f"{BASE}/auth")
        try:
            while ctx.pages:
                ctx.storage_state(path=str(STATE))
                time.sleep(2)
        except Exception:
            pass
    print("Sesion guardada." if STATE.exists() else "No se guardo sesion.")


def opciones_sesion():
    """Restaura la sesion guardada (cookies + localStorage) si existe."""
    import json
    if not STATE.exists():
        return {}
    st = json.loads(STATE.read_text(encoding="utf8"))
    ls = {}
    for o in st.get("origins", []):
        if "soloduenos.com" in o["origin"]:
            ls = {i["name"]: i["value"] for i in o["localStorage"]}
    js = "try{const d=%s;for(const k in d)localStorage.setItem(k,d[k])}catch(e){}" % json.dumps(ls)
    return {"cookies": st.get("cookies", []), "init_script_text": js}


def cmd_run(a):
    wb, ws = abrir_excel(a.excel)
    hdr, idx = indice_por_link(ws)
    zona_f = norm(a.zona)
    hoy = datetime.date.today().isoformat()
    vistos = nuevos = con_tel = 0
    resultados = {}

    ses = opciones_sesion()
    kw = {"cookies": ses["cookies"]} if ses else {}
    if ses:
        tmp = PROFILE / "init.js"
        tmp.write_text(ses["init_script_text"], encoding="utf8")
        kw["init_script"] = str(tmp)
    with DynamicSession(headless=a.headless, network_idle=True, **kw) as s:
        for op in a.ops:
            for url in listar_links(s, op, a.max_pages):
                vistos += 1
                if url in idx and ws.cell(idx[url], hdr.index("Teléfono / WhatsApp") + 1).value:
                    continue
                out = {"contacto": "", "login": False}
                resp = s.fetch(url, wait_selector="button:has-text('Contactar')", page_action=accion_detalle(not a.sin_contacto, out))
                datos = parsear(url, resp.get_all_text(separator="\n"))
                datos["contacto"] = out["contacto"]
                if zona_f and zona_f not in norm(datos["zona"]) and not (zona_f == "capital federal" and "ciudad autonoma" in norm(datos["zona"])):
                    continue
                resultados[url] = datos
                time.sleep(1)

    for url, d in resultados.items():
        tel, mail = parsear_contacto(d.get("contacto", ""), "")
        fila = {
            "ID": d["id"], "Fuente": FUENTE, "Link publicación": url, "Fecha de captura": hoy,
            "Nombre del dueño": d["nombre"], "Teléfono / WhatsApp": tel, "Email": mail, "Otro contacto": "",
            "Tipo": d["tipo"], "Operación": d["operacion"], "Zona": d["zona"], "Precio y moneda": d["precio"],
            "Características": d["caracteristicas"], "Detalle breve": detalle_breve(d), "Días publicado": "",
            "Estado de contacto": "nuevo", "Fecha de contacto": "", "Notas": "",
        }
        if url in idx:
            r = idx[url]
            for k in ("Nombre del dueño", "Teléfono / WhatsApp", "Email"):
                c = hdr.index(k) + 1
                if not ws.cell(r, c).value and fila[k]:
                    ws.cell(r, c).value = fila[k]
        else:
            ws.append([fila[c] for c in COLUMNAS])
            nuevos += 1
        con_tel += 1 if tel else 0

    wb["Corridas"].append([hoy, ",".join(a.ops), a.zona, vistos, nuevos, con_tel])
    wb.save(a.excel)
    print(f"Avisos vistos: {vistos} | en zona: {len(resultados)} | nuevos: {nuevos} | con teléfono: {con_tel}")
    if con_tel == 0 and not a.sin_contacto:
        print("AVISO: ningún teléfono. Si el sitio pide cuenta, corré 'login' primero.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("login").set_defaults(fn=cmd_login)
    r = sub.add_parser("run")
    r.add_argument("--excel", required=True)
    r.add_argument("--max-pages", type=int, default=3)
    r.add_argument("--ops", nargs="+", default=["venta", "alquiler"])
    r.add_argument("--zona", default="capital federal")
    r.add_argument("--sin-contacto", action="store_true")
    r.add_argument("--headless", action="store_true")
    r.set_defaults(fn=cmd_run)
    args = ap.parse_args()
    args.fn(args)
