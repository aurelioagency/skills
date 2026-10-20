"""Agrega leads nuevos (CSV) AL FINAL de la hoja Leads del Excel de campaña, sin tocar ninguna fila existente.

SOLO se corre cuando el usuario lo pidió ("cargalos al Excel"). Uso:
  python cargar_excel.py --excel "<Excel>" --csv tanda1.csv [--csv tanda2.csv ...] [--hoja Leads] [--rubro <id|nombre>]
         [--carga "2026-10-20 (búsqueda MX/CO)"] [--prueba] [--backup "<ruta del backup>"]
  --rubro rellena la columna "Rubro" de las filas nuevas (si la hoja no la tiene, la agrega al final). Si el CSV ya trae la columna
  Rubro, se usa esa.

Cómo trabaja (para no romper el archivo de nadie):
  1. copia el Excel a una carpeta temporal y trabaja ahí
  2. agrega las filas nuevas con los estados en NO (Mail enviado, IG enviado, WhatsApp enviado), la fórmula de "Para hacer"
     copiada de la última fila, el ID siguiente y el estilo de las filas anteriores; extiende filtro, listas desplegables y formato condicional
  3. comprueba que TODAS las celdas existentes de TODAS las hojas quedaron idénticas
  4. recién ahí reemplaza el original, y solo si nadie lo cambió mientras tanto (hash). Si el Excel está abierto en alguien
     (Windows lo bloquea), no pisa nada: avisa que hay que cerrarlo y se vuelve a correr.
No deja copias ni backups en las carpetas del usuario, salvo que se pase --backup. --prueba hace todo menos reemplazar el original.
El Excel de campaña tiene SOLO dos pestañas (Leads y Seguimientos): este script nunca crea pestañas.
"""
import copy
import hashlib
import os
import re
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from leads_common import *  # noqa: F401,F403,E402
import leads_common as C  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.formatting.formatting import ConditionalFormattingList  # noqa: E402
from openpyxl.worksheet.cell_range import MultiCellRange  # noqa: E402


def main():
    excel = C.arg_value("--excel")
    hoja = C.arg_value("--hoja", "Leads")
    carga = C.arg_value("--carga", "")
    backup = C.arg_value("--backup")
    prueba = "--prueba" in sys.argv
    rubro_arg = C.arg_value("--rubro")
    rubro_nombre = ""
    if rubro_arg:
        for rr in C.lista_rubros():
            if rubro_arg.lower() in (rr["id"].lower(), rr["nombre"].lower()):
                rubro_nombre = rr["nombre"]
        rubro_nombre = rubro_nombre or rubro_arg
    H = lambda p: hashlib.md5(open(p, "rb").read()).hexdigest()  # noqa: E731
    hash0 = H(excel)
    tmpdir = tempfile.mkdtemp(prefix="leads_")
    tmp = os.path.join(tmpdir, os.path.basename(excel))
    shutil.copy2(excel, tmp)
    if backup:
        shutil.copy2(excel, backup)
        print("backup en:", backup)
    xl = excel.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(tmp, keep_vba=xl)
    ws = wb[hoja]
    h = {str(c.value): i + 1 for i, c in enumerate(ws[1]) if c.value}
    n0 = ws.max_row
    snap = {s.title: [[c.value for c in row] for row in s.iter_rows()] for s in wb.worksheets}

    def col(*names):
        for n in names:
            for k, i in h.items():
                if k.lower() == n.lower():
                    return i
        return None

    c_emp, c_pais = col("Empresa"), col("País", "Pais")
    c_mail, c_ig, c_wa = col("Mail enviado", "Mail_Enviado"), col("IG enviado", "IG_Enviado"), col("WhatsApp enviado", "WhatsApp_Enviado")
    c_tel, c_email, c_insta = col("Teléfono / WhatsApp", "Telefono / WhatsApp"), col("Email"), col("Instagram")
    c_id, c_para = col("ID"), col("Para hacer")
    c_web, c_ml, c_nota, c_carga = col("Web"), col("Tienda Mercado Libre", "Tienda_MercadoLibre"), col("Notas y fuente de los datos"), col("Carga")
    c_rubro = col("Rubro")
    faltan = [n for n, c in (("Empresa", c_emp), ("País", c_pais), ("Instagram", c_insta), ("Email", c_email)) if c is None]
    if faltan:
        sys.exit(f"ABORTO: la hoja {hoja} no tiene las columnas {faltan}")
    idx = C.build_index(tmp, hoja)
    rows = []
    for p in C.arg_list("--csv"):
        rows += [(p, r) for r in csv.DictReader(open(p, encoding="utf-8-sig"))]
    nuevos, omit = [], []
    for p, r in rows:
        hd = C.handle_of(r["Instagram"])
        ph = {C.d10(x) for x in C.split(r["WhatsApp"]) + C.split(r["Telefono"])}
        if C.norm(r["Empresa"]) in idx["names"] or (hd and hd in idx["handles"]) or (ph & idx["phones"]):
            omit.append(r["Empresa"])
            continue
        idx["names"].add(C.norm(r["Empresa"]))
        if hd:
            idx["handles"].add(hd)
        idx["phones"] |= ph
        nuevos.append(r)
    print("a cargar:", len(nuevos), "| omitidos porque ya existen:", len(omit), omit[:5])
    if c_rubro is None and (rubro_nombre or any((r.get("Rubro") or "").strip() for r in nuevos)):
        c_rubro = ws.max_column + 1
        ws.cell(1, c_rubro, "Rubro")
        ws.cell(1, c_rubro)._style = copy.copy(ws.cell(1, c_rubro - 1)._style)
        print("la hoja no tenía la columna Rubro: se agregó al final")
    last = n0
    while last > 1 and not ws.cell(last, c_emp).value:
        last -= 1
    ids = [ws.cell(r, c_id).value for r in range(2, last + 1) if c_id and isinstance(ws.cell(r, c_id).value, (int, float))]
    maxid = max(ids) if ids else 0
    formula = ws.cell(last, c_para).value if c_para and isinstance(ws.cell(last, c_para).value, str) and ws.cell(last, c_para).value.startswith("=") else None
    ncol = ws.max_column
    for i, r in enumerate(nuevos):
        rl = last + 1 + i
        src = rl - 2 if rl - 2 > 1 else rl - 1
        for c in range(1, ncol + 1):
            s = ws.cell(src, c)
            if s.has_style:
                ws.cell(rl, c)._style = copy.copy(s._style)
        wa_nums, tel_nums = C.split(r["WhatsApp"]), C.split(r["Telefono"])
        ws.cell(rl, c_emp, r["Empresa"])
        ws.cell(rl, c_pais, r["Pais"])
        if formula and c_para:
            ws.cell(rl, c_para, re.sub(r"([A-Z]{1,2})%d\b" % last, r"\g<1>%d" % rl, formula))
        for c in (c_mail, c_ig, c_wa):
            if c:
                ws.cell(rl, c, "NO")
        # Teléfono / WhatsApp: los de WhatsApp primero; si no hay, los teléfonos encontrados (no se sacan números por suposición)
        contacto = "; ".join(wa_nums) if wa_nums else "; ".join(tel_nums)
        if contacto and c_tel:
            ws.cell(rl, c_tel, contacto)
        if r["Email"]:
            ws.cell(rl, c_email, r["Email"])
        if r["Instagram"]:
            ws.cell(rl, c_insta, r["Instagram"])
        if c_id:
            ws.cell(rl, c_id, maxid + 1 + i)
        if c_web and r["Web"]:
            ws.cell(rl, c_web, r["Web"])
        if c_ml and r.get("Tienda_MercadoLibre"):
            ws.cell(rl, c_ml, r["Tienda_MercadoLibre"])
        if c_nota:
            extra = f" | Teléfonos que no son WhatsApp confirmado: {'; '.join(tel_nums)}" if wa_nums and tel_nums else ""
            ws.cell(rl, c_nota, f"{r['Estado']}. Fuente de cada dato: {r['De_donde_sale']}{extra}")
        if c_carga and carga:
            ws.cell(rl, c_carga, carga)
        rb = (r.get("Rubro") or "").strip() or rubro_nombre
        if c_rubro and rb:
            ws.cell(rl, c_rubro, rb)
    nl = last + len(nuevos)
    if nuevos:
        if ws.auto_filter.ref:
            ws.auto_filter.ref = re.sub(r"\d+$", str(nl), ws.auto_filter.ref)
        for dv in ws.data_validations.dataValidation:
            dv.sqref = MultiCellRange(re.sub(str(n0), str(nl), str(dv.sqref)))
        newcf = ConditionalFormattingList()
        for cf in ws.conditional_formatting:
            rng = re.sub(str(n0), str(nl), str(cf.sqref))
            for rule in cf.rules:
                newcf.add(rng, rule)
        ws.conditional_formatting = newcf
    wb.save(tmp)
    # ---- verificación
    w2 = openpyxl.load_workbook(tmp, keep_vba=xl)
    ok = True
    for t, v in snap.items():
        n = len(v) if t != hoja else n0
        now = [[c.value for c in row] for row in w2[t].iter_rows(min_row=1, max_row=n, max_col=len(v[0]) if v else None)]
        if now != [row[: len(now[0])] if now else row for row in v[:n]]:
            ok = False
            print("CAMBIÓ algo existente en la hoja", t)
    print("hojas:", w2.sheetnames, "| filas en", hoja, ":", w2[hoja].max_row, "| existentes intactas:", ok)
    if not ok:
        sys.exit("ABORTO: no se pisa el original")
    if H(excel) != hash0:
        sys.exit("ABORTO: el Excel cambió mientras trabajaba (otra persona o sesión lo guardó). Volver a correr.")
    if prueba or not nuevos:
        print("PRUEBA: no se tocó el Excel real" if prueba else "Nada para cargar")
        shutil.rmtree(tmpdir, ignore_errors=True)
        return
    try:
        shutil.copy2(tmp, excel)
    except PermissionError:
        shutil.rmtree(tmpdir, ignore_errors=True)
        sys.exit("El Excel está abierto en Excel (Windows lo bloquea). Pedirle al usuario que lo cierre y volver a correr. No se tocó nada.")
    shutil.rmtree(tmpdir, ignore_errors=True)
    print("EXCEL ACTUALIZADO:", len(nuevos), "filas nuevas al final de", hoja)


if __name__ == "__main__":
    main()
